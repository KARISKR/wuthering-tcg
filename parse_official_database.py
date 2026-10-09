import openpyxl
import json
import os
import re

# ==========================================
# 1. card_database.xlsx 파싱
# ==========================================
wb = openpyxl.load_workbook(r'c:\Users\Lyicos\.gemini\antigravity\scratch\wuthering-tcg\public\cards\card_database.xlsx', data_only=True)
ws = wb.active

headers = [c.value for c in ws[1]]
print("Headers:", headers)

# 컬럼명 매핑
# 'ID', '카드코드', '카드명', '카드타입', '효과/설명', '이미지URL', '로컬이미지경로',
# 'type_id', 'card_type', 'type_name', 'weapon_type_name', 'force_name', 'attr_name',
# 'feature_name', 'character_name', 'rarity_name', 'fee', 'level', 'speed', 'damage',
# 'obtain', 'info', 'color_name'

element_map = {
    '회절': 'SPECTRO',
    '인멸': 'HAVOC',
    '기류': 'AERO',
    '용융': 'FUSION',
    '응결': 'GLACIO',
}

color_map = {
    '적색': 'RED',
    '녹색': 'GREEN',
    '청색': 'BLUE',
    'RED': 'RED',
    'GREEN': 'GREEN',
    'BLUE': 'BLUE',
}

cards_by_code = {}

for row in ws.iter_rows(min_row=2, values_only=True):
    d = dict(zip(headers, row))
    code = d.get('카드코드')
    if not code:
        continue
    
    # 중복 카드코드 처리: 이미 존재하는 경우 일러스트/레어도 패러렐 버전이거나 중복
    # 대표 카드로 저장
    if code in cards_by_code:
        continue
        
    card_type_raw = str(d.get('card_type') or '').strip().lower()
    type_name_raw = str(d.get('type_name') or '').strip()
    is_char = (card_type_raw == 'character') or ('캐릭터' in type_name_raw)
    
    name = str(d.get('카드명') or '').strip()
    char_name_field = str(d.get('character_name') or '').strip()
    if char_name_field == '-': char_name_field = ''
    
    # 속성 매핑
    attr_raw = str(d.get('attr_name') or '').strip()
    elem = element_map.get(attr_raw, 'SPECTRO')
    
    # 색상 매핑
    col_raw = str(d.get('color_name') or '').strip()
    col = color_map.get(col_raw, None)
    
    # 레벨
    lvl_val = d.get('level')
    level = None
    if lvl_val is not None and str(lvl_val).strip().isdigit():
        level = int(str(lvl_val).strip())
    elif is_char:
        level = 0
        
    # 비용 (fee)
    fee_val = d.get('fee')
    cost = 0
    if fee_val is not None and str(fee_val).strip().isdigit():
        cost = int(str(fee_val).strip())
        
    # 속도 (speed)
    speed_val = d.get('speed')
    speed = None
    if speed_val is not None and str(speed_val).strip().isdigit():
        speed = int(str(speed_val).strip())
        
    # 데미지 (damage)
    dmg_val = d.get('damage')
    damage = 0
    if dmg_val is not None and str(dmg_val).strip().isdigit():
        damage = int(str(dmg_val).strip())
        
    # 공식 효과 텍스트 (info)
    info_text = str(d.get('info') or '').strip()
    if not info_text or info_text == 'None':
        info_text = str(d.get('효과/설명') or '').strip()
    if not info_text or info_text == 'None':
        info_text = f"공식 {type_name_raw} [{name}]."
        
    # 로컬 이미지 경로
    local_img = str(d.get('로컬이미지경로') or '').strip().replace('\\', '/')
    if local_img and not local_img.startswith('/'):
        local_img = f"/cards/{local_img}"
        
    # 추격 수치 파싱: info에 '추격 X' 또는 '추격X'가 있으면 추출
    pursuit_cnt = None
    m_pur = re.search(r'추격\s*(\d+)', info_text)
    if m_pur:
        pursuit_cnt = int(m_pur.group(1))
        
    card_obj = {
        'id': f"card-{code}",
        'code': code,
        'nameKr': name,
        'kind': 'CHARACTER' if is_char else 'ACTION',
        'type_name': type_name_raw,
        'characterName': char_name_field or (name.split('(')[0] if is_char else ''),
        'level': level,
        'element': elem,
        'weaponType': str(d.get('weapon_type_name') or '-').strip(),
        'color': col if not is_char else None,
        'cost': cost,
        'speed': speed,
        'damage': damage,
        'pursuitCount': pursuit_cnt,
        'characterExclusive': char_name_field if (not is_char and char_name_field) else None,
        'description': info_text,
        'artUrl': local_img,
        'feature': str(d.get('feature_name') or '-').strip(),
        'rarity': str(d.get('rarity_name') or '★').strip(),
        'obtain': str(d.get('obtain') or '').strip(),
    }
    
    cards_by_code[code] = card_obj

print(f"Total parsed unique cards: {len(cards_by_code)}")
char_count = sum(1 for c in cards_by_code.values() if c['kind'] == 'CHARACTER')
action_count = sum(1 for c in cards_by_code.values() if c['kind'] == 'ACTION')
print(f"Character cards: {char_count}, Action cards: {action_count}")

# JSON으로 저장
with open(r'c:\Users\Lyicos\.gemini\antigravity\scratch\wuthering-tcg\src\data\official_database.json', 'w', encoding='utf-8') as f:
    json.dump(list(cards_by_code.values()), f, ensure_ascii=False, indent=2)

print("Saved to src/data/official_database.json")
