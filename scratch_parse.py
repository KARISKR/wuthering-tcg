import openpyxl
import json
import os
from pypdf import PdfReader

# 1. Excel 분석
wb = openpyxl.load_workbook(r'c:\Users\Lyicos\.gemini\antigravity\scratch\wuthering-tcg\public\cards\card_database.xlsx', data_only=True)
ws = wb.active

headers = [cell.value for cell in ws[1]]
print("Headers:", headers)

cards = []
for row_idx in range(2, ws.max_row + 1):
    row_vals = [cell.value for cell in ws[row_idx]]
    if not any(row_vals):
        continue
    row_dict = {}
    for h, val in zip(headers, row_vals):
        row_dict[h] = val
    cards.append(row_dict)

print(f"Total cards in Excel: {len(cards)}")
# 샘플 3개 출력
for c in cards[:3]:
    print("Sample card:", json.dumps(c, ensure_ascii=False))

# 저장
with open(r'c:\Users\Lyicos\.gemini\antigravity\scratch\wuthering-tcg\src\data\extracted_cards.json', 'w', encoding='utf-8') as f:
    json.dump(cards, f, ensure_ascii=False, indent=2)

# 2. PDF 분석
rules_dir = r'c:\Users\Lyicos\.gemini\antigravity\scratch\wuthering-tcg\rules'
pdf_files = [f for f in os.listdir(rules_dir) if f.endswith('.pdf')]
print("\nPDF files found:", pdf_files)

for pdf_name in pdf_files:
    pdf_path = os.path.join(rules_dir, pdf_name)
    try:
        reader = PdfReader(pdf_path)
        print(f"\n--- PDF: {pdf_name} (Pages: {len(reader.pages)}) ---")
        text_preview = ""
        for page_idx in range(min(5, len(reader.pages))):
            text_preview += f"\n[Page {page_idx+1}]\n" + (reader.pages[page_idx].extract_text() or "")
        
        txt_out = os.path.join(rules_dir, pdf_name.replace('.pdf', '.txt'))
        full_text = ""
        for page_idx, p in enumerate(reader.pages):
            full_text += f"\n=== [Page {page_idx+1}] ===\n" + (p.extract_text() or "")
        with open(txt_out, 'w', encoding='utf-8') as tf:
            tf.write(full_text)
        print(f"Extracted to {txt_out} (Chars: {len(full_text)})")
        print("First 300 chars preview:\n", full_text[:300].strip())
    except Exception as e:
        print(f"Error reading {pdf_name}: {e}")
